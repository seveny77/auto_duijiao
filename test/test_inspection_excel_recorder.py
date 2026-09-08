# -*- coding: utf-8 -*-
"""单工作表正式检测 Excel 记录测试。"""

from datetime import datetime
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from backend.inspection_excel_recorder import HEADERS, InspectionExcelRecorder
from backend.inspection_types import (
    CircleCandidate,
    CircleInspectionResult,
    DefectMeasurement,
    ImageInspectionResult,
    InspectionStatus,
    SizeRuleInspectionResult,
)


def _circle(status, spot_count, scratch_count, *, found=True):
    measurements = []
    for index in range(spot_count):
        measurements.append(DefectMeasurement(
            instance_index=index,
            class_id=0,
            class_name="污点",
            area_um2=1.0,
            source="polygon",
        ))
    for index in range(scratch_count):
        measurements.append(DefectMeasurement(
            instance_index=spot_count + index,
            class_id=1,
            class_name="划痕",
            area_um2=2.0,
            source="polygon",
        ))
    return CircleInspectionResult(
        circle_id="circle-001",
        circle_candidate=(
            CircleCandidate(center_x=10, center_y=10) if found else None
        ),
        completed=found,
        circle_confirmed=found,
        status=status,
        measurements=measurements if found else [],
        size_rule_results=(
            [
                SizeRuleInspectionResult(
                    rule_id="spot",
                    defect_class="污点",
                    actual_instance_count=spot_count,
                ),
                SizeRuleInspectionResult(
                    rule_id="scratch",
                    defect_class="划痕",
                    actual_instance_count=scratch_count,
                ),
            ]
            if found else []
        ),
    )


class InspectionExcelRecorderTest(unittest.TestCase):
    def test_single_and_five_face_results_append_to_one_sheet(self):
        timestamp = datetime(2026, 9, 7, 14, 30, 15, 123000)
        single = ImageInspectionResult(
            expected_circle_count=1,
            detected_circle_count=1,
            status=InspectionStatus.FAIL,
            circle_results=[_circle(InspectionStatus.FAIL, 1, 2)],
            failure_reasons=["circle-001: 污点尺寸段数量超过上限"],
        )
        five = ImageInspectionResult(
            expected_circle_count=5,
            detected_circle_count=4,
            status=InspectionStatus.PENDING,
            circle_results=[
                _circle(InspectionStatus.PASS, 0, 0),
                _circle(InspectionStatus.FAIL, 2, 1),
                _circle(InspectionStatus.PASS, 3, 0),
                _circle(InspectionStatus.PASS, 0, 4),
                _circle(InspectionStatus.PENDING, 0, 0, found=False),
            ],
            warnings=["circle-005 未找到对应候选圆"],
        )

        with tempfile.TemporaryDirectory() as directory:
            recorder = InspectionExcelRecorder(directory)
            path = recorder.append(
                single,
                completed_at=timestamp,
                material_number="M-001",
                original_image_path=str(Path(directory) / "single.jpg"),
                record_id="SINGLE-001",
            )
            recorder.append(
                five,
                completed_at=timestamp,
                material_number="M-005",
                original_image_path=str(Path(directory) / "five.jpg"),
                record_id="FIVE-001",
            )

            workbook = load_workbook(path, data_only=True)
            self.assertEqual(workbook.sheetnames, ["检测记录"])
            sheet = workbook["检测记录"]
            self.assertEqual(tuple(cell.value for cell in sheet[1]), HEADERS)
            self.assertEqual(sheet.max_row, 3)

            self.assertEqual(sheet.cell(2, 1).value, "SINGLE-001")
            self.assertEqual(sheet.cell(2, 3).value, "单圆物料")
            self.assertEqual(sheet.cell(2, 4).value, "M-001")
            self.assertEqual(sheet.cell(2, 7).value, "不合格")
            self.assertEqual(sheet.cell(2, 8).value, 1)
            self.assertEqual(sheet.cell(2, 9).value, 2)
            self.assertTrue(all(
                sheet.cell(2, column).value is None
                for column in range(10, 18)
            ))
            self.assertEqual(sheet.cell(2, 18).value, 3)
            self.assertEqual(sheet.cell(2, 19).value, 5.0)

            self.assertEqual(sheet.cell(3, 3).value, "五圆物料")
            self.assertEqual(sheet.cell(3, 4).value, "M-005")
            self.assertEqual(sheet.cell(3, 6).value, 4)
            self.assertEqual(sheet.cell(3, 7).value, "待确认")
            self.assertEqual(
                [sheet.cell(3, column).value for column in range(8, 16)],
                [0, 0, 2, 1, 3, 0, 0, 4],
            )
            self.assertIsNone(sheet.cell(3, 16).value)
            self.assertIsNone(sheet.cell(3, 17).value)
            self.assertEqual(sheet.cell(3, 20).value, 3)
            self.assertEqual(sheet.cell(3, 21).value, 1)
            self.assertEqual(sheet.cell(3, 22).value, 1)
            self.assertIn("端面5未找到", sheet.cell(3, 23).value)
            workbook.close()

            self.assertFalse(list(Path(directory).rglob("*.tmp.xlsx")))


if __name__ == "__main__":
    unittest.main()
