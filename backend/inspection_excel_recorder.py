# -*- coding: utf-8 -*-
"""把一次正式质检结果追加到按日保存的单工作表 Excel。"""

from datetime import datetime
import os
from pathlib import Path
import threading
import uuid

from backend.inspection_types import (
    DefectMeasurement,
    ImageInspectionResult,
    InspectionStatus,
)


SHEET_NAME = "检测记录"
MAX_FACE_COUNT = 5
HEADERS = (
    "检测ID",
    "检测时间",
    "物料类型",
    "物料编号",
    "预期端面数",
    "实际检出端面数",
    "总判定",
    "端面1污点数量",
    "端面1划痕数量",
    "端面2污点数量",
    "端面2划痕数量",
    "端面3污点数量",
    "端面3划痕数量",
    "端面4污点数量",
    "端面4划痕数量",
    "端面5污点数量",
    "端面5划痕数量",
    "有效缺陷总数",
    "缺陷总面积（um²）",
    "合格端面数",
    "不合格端面数",
    "待确认端面数",
    "失败摘要",
    "原图路径",
    "结果图路径",
    "备注",
)

_STATUS_TEXT = {
    InspectionStatus.PASS: "合格",
    InspectionStatus.FAIL: "不合格",
    InspectionStatus.PENDING: "待确认",
    InspectionStatus.ERROR: "检测错误",
}


class InspectionExcelRecorder:
    """以原子替换方式维护每日一个 ``.xlsx`` 检测记录文件。"""

    def __init__(self, root: str):
        root_text = str(root or "").strip()
        if not root_text:
            raise ValueError("Excel 检测记录目录不能为空")
        self._root = os.path.abspath(root_text)
        self._lock = threading.Lock()

    @property
    def root(self) -> str:
        return self._root

    def append(
        self,
        result: ImageInspectionResult,
        *,
        completed_at: datetime | None = None,
        material_number: str = "",
        original_image_path: str = "",
        result_image_path: str = "",
        remark: str = "",
        record_id: str = "",
    ) -> str:
        """追加一行并返回当日 Excel 的绝对路径。"""

        if not isinstance(result, ImageInspectionResult):
            raise TypeError("result 必须是 ImageInspectionResult")
        if result.expected_circle_count > MAX_FACE_COUNT:
            raise ValueError("Excel 检测记录当前最多支持 5 个端面")

        timestamp = completed_at or datetime.now()
        output_path = self._daily_path(timestamp)
        row = build_record_row(
            result,
            completed_at=timestamp,
            material_number=material_number,
            original_image_path=original_image_path,
            result_image_path=result_image_path,
            remark=remark,
            record_id=record_id,
        )
        with self._lock:
            self._append_atomic(output_path, row)
        return str(output_path)

    def _daily_path(self, timestamp: datetime) -> Path:
        month_dir = Path(self._root) / timestamp.strftime("%Y-%m")
        # 新规则使用“污点/划痕 + 尺寸段”结果，和旧版“脏污/异物”
        # 的表头不同。使用 v2 文件名，保留旧表格，避免新旧格式混写。
        return month_dir / f"检测记录_v2_{timestamp:%Y%m%d}.xlsx"

    @staticmethod
    def _append_atomic(output_path: Path, row: list):
        # 生产运行时才导入，软件启动阶段不增加 Excel 模块加载开销。
        from openpyxl import Workbook, load_workbook

        output_path.parent.mkdir(parents=True, exist_ok=True)
        if output_path.exists():
            workbook = load_workbook(output_path)
            _validate_workbook(workbook)
            worksheet = workbook[SHEET_NAME]
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = SHEET_NAME
            _initialize_worksheet(worksheet)

        worksheet.append(row)
        _format_appended_row(worksheet, worksheet.max_row)
        worksheet.auto_filter.ref = (
            f"A1:{_column_letter(len(HEADERS))}{worksheet.max_row}"
        )
        temporary_path = output_path.with_name(
            f".{output_path.stem}.{uuid.uuid4().hex}.tmp.xlsx"
        )
        try:
            workbook.save(temporary_path)
            workbook.close()
            os.replace(temporary_path, output_path)
        finally:
            try:
                workbook.close()
            finally:
                if temporary_path.exists():
                    temporary_path.unlink()


def build_record_row(
    result: ImageInspectionResult,
    *,
    completed_at: datetime,
    material_number: str = "",
    original_image_path: str = "",
    result_image_path: str = "",
    remark: str = "",
    record_id: str = "",
) -> list:
    """把多端面结果展开成固定 26 列的一行。

    端面数量列来自新尺寸规则引擎的 ``size_rule_results``；判定结果
    仍以端面状态和 ``failure_reasons`` 为准，不由 Excel 表格重复判定。
    """

    expected_count = max(0, int(result.expected_circle_count))
    if expected_count > MAX_FACE_COUNT:
        raise ValueError("Excel 检测记录当前最多支持 5 个端面")

    face_counts: list[int | None] = []
    for face_index in range(MAX_FACE_COUNT):
        if face_index >= expected_count:
            face_counts.extend((None, None))
            continue
        circle_result = (
            result.circle_results[face_index]
            if face_index < len(result.circle_results)
            else None
        )
        face_counts.extend((
            _class_valid_count(circle_result, "污点"),
            _class_valid_count(circle_result, "划痕"),
        ))

    expected_results = list(result.circle_results[:expected_count])
    passed_count = sum(
        item.status == InspectionStatus.PASS for item in expected_results
    )
    failed_count = sum(
        item.status == InspectionStatus.FAIL for item in expected_results
    )
    pending_count = max(0, expected_count - passed_count - failed_count)
    recordable_results = [
        item for item in expected_results if _is_recordable_circle(item)
    ]
    total_count = (
        sum(_circle_valid_count(item) for item in recordable_results)
        if recordable_results else None
    )
    total_area = (
        sum(_circle_total_area_um2(item) for item in recordable_results)
        if recordable_results else None
    )

    original_path = str(original_image_path or "")
    overlay_path = str(result_image_path or "")
    if not overlay_path and original_path:
        source = Path(original_path)
        overlay_path = str(source.with_name(f"{source.stem}_inspection.jpg"))

    row_record_id = str(record_id or "").strip() or _new_record_id(completed_at)
    material_type = (
        "单圆物料" if expected_count == 1
        else "五圆物料" if expected_count == 5
        else f"{expected_count}圆物料"
    )
    return [
        row_record_id,
        completed_at,
        material_type,
        str(material_number or ""),
        expected_count,
        int(result.detected_circle_count),
        _status_text(result.status),
        *face_counts,
        total_count,
        total_area,
        passed_count,
        failed_count,
        pending_count,
        _failure_summary(result),
        original_path,
        overlay_path,
        str(remark or "") or _warning_summary(result),
    ]


def _is_recordable_circle(circle_result) -> bool:
    return bool(
        circle_result is not None
        and circle_result.circle_candidate is not None
        and circle_result.completed
        and circle_result.circle_confirmed
        and circle_result.status in (InspectionStatus.PASS, InspectionStatus.FAIL)
    )


def _class_valid_count(circle_result, class_name: str) -> int | None:
    """返回新规则引擎确认的类别数量。

    空白表示该端面没有完成正式判定，0 表示已判定且没有该类缺陷。
    同一个实例只会命中一个尺寸段，因此对规则结果求和不会重复计数。
    """

    if not _is_recordable_circle(circle_result):
        return None
    normalized_name = str(class_name).strip().casefold()
    size_rule_results = list(circle_result.size_rule_results or [])
    if size_rule_results:
        return sum(
            max(0, int(item.actual_instance_count))
            for item in size_rule_results
            if str(item.defect_class).strip().casefold() == normalized_name
        )

    # 允许对仅携带测量结果的中间快照生成记录；正式检测结果通常会同时
    # 携带 size_rule_results。
    return sum(
        1 for item in _valid_measurements(circle_result)
        if str(item.class_name).strip().casefold() == normalized_name
    )


def _circle_valid_count(circle_result) -> int:
    return sum(
        _class_valid_count(circle_result, class_name) or 0
        for class_name in ("污点", "划痕")
    )


def _circle_total_area_um2(circle_result) -> float:
    """按每个实例的物理面积汇总，避免旧 mm² 区域字段混入新表。"""

    return sum(
        max(0.0, float(item.area_um2))
        for item in _valid_measurements(circle_result)
        if item.area_um2 is not None
    )


def _valid_measurements(circle_result) -> list[DefectMeasurement]:
    return [
        item for item in list(circle_result.measurements or [])
        if isinstance(item, DefectMeasurement)
        and str(item.class_name).strip() in {"污点", "划痕"}
        and item.source != "invalid"
    ]


def _status_text(status) -> str:
    try:
        normalized = (
            status if isinstance(status, InspectionStatus)
            else InspectionStatus(status)
        )
    except (TypeError, ValueError):
        return "待确认"
    return _STATUS_TEXT[normalized]


def _failure_summary(result: ImageInspectionResult) -> str:
    values = list(result.failure_reasons or [])
    if result.error:
        values.append(str(result.error))
    expected_count = max(0, int(result.expected_circle_count))
    for index in range(expected_count):
        circle = (
            result.circle_results[index]
            if index < len(result.circle_results) else None
        )
        if circle is None or circle.circle_candidate is None:
            values.append(f"端面{index + 1}未找到")
        elif circle.status == InspectionStatus.ERROR:
            values.append(f"端面{index + 1}检测错误")
        elif circle.status == InspectionStatus.PENDING:
            values.append(f"端面{index + 1}待确认")
    return _unique_join(values)


def _warning_summary(result: ImageInspectionResult) -> str:
    return _unique_join(result.warnings or [])


def _unique_join(values) -> str:
    unique_values = []
    for value in values:
        text = str(value or "").strip()
        if text and text not in unique_values:
            unique_values.append(text)
    return "；".join(unique_values)


def _new_record_id(timestamp: datetime) -> str:
    prefix = timestamp.strftime("%Y%m%d_%H%M%S_%f")[:-3]
    return f"{prefix}_{uuid.uuid4().hex[:4].upper()}"


def _initialize_worksheet(worksheet):
    from openpyxl.styles import Alignment, Font, PatternFill

    worksheet.append(list(HEADERS))
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False
    worksheet.row_dimensions[1].height = 32
    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(name="Microsoft YaHei", color="FFFFFF", bold=True)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    widths = {
        "A": 25, "B": 23, "C": 12, "D": 20, "E": 12, "F": 14,
        "G": 12, "R": 14, "S": 18, "T": 14, "U": 14, "V": 14,
        "W": 48, "X": 48, "Y": 48, "Z": 20,
    }
    for index in range(8, 18):
        widths[_column_letter(index)] = 16
    for column, width in widths.items():
        worksheet.column_dimensions[column].width = width


def _format_appended_row(worksheet, row_index: int):
    from openpyxl.styles import Alignment, Font, PatternFill

    status_fills = {
        "合格": "C6EFCE",
        "不合格": "FFC7CE",
        "待确认": "FFEB9C",
        "检测错误": "F4CCCC",
    }
    for cell in worksheet[row_index]:
        cell.font = Font(name="Microsoft YaHei", size=10)
        cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[row_index].height = 22
    worksheet.cell(row_index, 2).number_format = "yyyy-mm-dd hh:mm:ss.000"
    worksheet.cell(row_index, 19).number_format = "0.00"
    status_cell = worksheet.cell(row_index, 7)
    status_cell.alignment = Alignment(horizontal="center", vertical="center")
    fill_color = status_fills.get(status_cell.value)
    if fill_color:
        status_cell.fill = PatternFill("solid", fgColor=fill_color)
    for column_index in (24, 25):
        cell = worksheet.cell(row_index, column_index)
        if cell.value:
            cell.hyperlink = _file_uri(str(cell.value))
            cell.style = "Hyperlink"


def _validate_workbook(workbook):
    if workbook.sheetnames != [SHEET_NAME]:
        raise ValueError("检测记录 Excel 必须且只能包含“检测记录”工作表")
    worksheet = workbook[SHEET_NAME]
    actual = tuple(
        worksheet.cell(1, index).value for index in range(1, len(HEADERS) + 1)
    )
    if actual != HEADERS:
        raise ValueError("检测记录 Excel 表头与当前版本不一致")


def _file_uri(path: str) -> str:
    return Path(os.path.abspath(path)).as_uri()


def _column_letter(index: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(index)
